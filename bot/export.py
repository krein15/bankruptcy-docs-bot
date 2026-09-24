"""Excel-отчёт для юриста: строка — клиент, столбец — документ, цвет — статус."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bot import texts
from bot.checklist import CHECKLIST, progress, status_of

FILLS = {
    "review": PatternFill("solid", fgColor="FFF2CC"),    # жёлтый — ждёт проверки
    "accepted": PatternFill("solid", fgColor="D9EAD3"),  # зелёный — принят
    "rejected": PatternFill("solid", fgColor="F4CCCC"),  # красный — вернули
    "na": PatternFill("solid", fgColor="EEEEEE"),        # серый — документа нет
}


def build_report(clients: list, all_statuses: dict) -> bytes:
    """Возвращает готовый .xlsx в памяти — на диск ничего не пишем."""
    wb = Workbook()
    ws = wb.active
    ws.title = texts.EXPORT_SHEET

    ws.append(texts.EXPORT_HEADERS + [item["title"] for item in CHECKLIST])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    first_doc_column = len(texts.EXPORT_HEADERS) + 1
    for client in clients:
        statuses = all_statuses.get(client["id"], {})
        done, total = progress(statuses)
        ws.append([client["full_name"], client["phone"], texts.fmt_time(client["last_activity_at"]), f"{done} из {total}"])
        row = ws.max_row
        for column, item in enumerate(CHECKLIST, start=first_doc_column):
            status = status_of(statuses, item["id"])
            cell = ws.cell(row=row, column=column, value=texts.STATUS_LABELS[status])
            if status in FILLS:
                cell.fill = FILLS[status]
            # Причина возврата — примечанием к ячейке: видно при наведении мыши
            if status == "rejected" and statuses[item["id"]]["comment"]:
                cell.comment = Comment(statuses[item["id"]]["comment"], texts.EXPORT_COMMENT_AUTHOR)

    widths = [32, 16, 20, 10] + [16] * len(CHECKLIST)
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = ws.cell(row=2, column=2)  # шапка и ФИО не уезжают при прокрутке
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
