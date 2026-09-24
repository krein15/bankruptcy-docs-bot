"""Выгрузка в Excel: открываем присланный файл и смотрим, что внутри."""
from io import BytesIO

from openpyxl import load_workbook

from bot import texts
from bot.checklist import CHECKLIST


def workbook_from(replies):
    [document] = replies.of("SendDocument")
    assert document.document.filename.startswith("Документы клиентов")
    assert document.document.filename.endswith(".xlsx")
    return load_workbook(BytesIO(document.document.data)).active


async def test_export(ivan, petr, lawyer_chat):
    await ivan.register()
    await ivan.submit("passport")
    await lawyer_chat.reject(ivan.id, "passport", "Размыто")
    await petr.register(full_name="Петров Пётр Петрович")
    await petr.submit("snils")
    await lawyer_chat.accept(petr.id, "snils")

    sheet = workbook_from(await lawyer_chat.send("/export"))
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:4] == tuple(texts.EXPORT_HEADERS)
    assert len(rows[0]) == len(texts.EXPORT_HEADERS) + len(CHECKLIST)
    assert len(rows) == 3, "шапка + 2 клиента"

    ivan_row = next(i for i, row in enumerate(rows, start=1) if row[0] == "Иванов Иван Иванович")
    passport = sheet.cell(row=ivan_row, column=5)
    assert passport.value == "вернули"
    assert passport.comment.text == "Размыто", "причина возврата — в примечании"
    assert passport.fill.fgColor.rgb.endswith("F4CCCC"), "красная заливка"

    petr_row = next(i for i, row in enumerate(rows, start=1) if row[0] == "Петров Пётр Петрович")
    assert sheet.cell(row=petr_row, column=6).value == "принят"
    assert sheet.cell(row=petr_row, column=4).value == f"1 из {len(CHECKLIST)}"
    assert sheet.freeze_panes == "B2", "шапка и ФИО закреплены"


async def test_export_button_same_as_command(ivan, lawyer_chat):
    await ivan.register()
    replies = await lawyer_chat.press("export")
    assert workbook_from(replies)


async def test_export_without_clients(lawyer_chat):
    replies = await lawyer_chat.send("/export")
    assert texts.NO_CLIENTS in replies.text
    assert not replies.of("SendDocument")


async def test_unfinished_registration_not_exported(ivan, lawyer_chat):
    await ivan.send("/start")
    replies = await lawyer_chat.send("/export")
    assert texts.NO_CLIENTS in replies.text
