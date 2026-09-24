"""Демо-режим: посетитель сам себе клиент и юрист, чужих данных не видит, владелец узнаёт о посетителях."""
import pytest

from bot import config, db, keyboards as kb, reminders, texts
from conftest import LAWYER_ID, LONG_AGO, set_client
from scripts import seed_demo


@pytest.fixture(autouse=True)
def demo(monkeypatch):
    monkeypatch.setattr(config, "DEMO_MODE", True)


async def test_intro_warns_not_to_send_real_documents(ivan):
    replies = await ivan.send("/start")
    assert "Не отправляйте настоящие документы" in replies.text
    assert "согласие" in replies.text


async def test_owner_learns_about_new_visitor(ivan):
    ivan.user["username"] = "ivan_lawyer"
    replies = await ivan.send("/start")
    note = replies.to(LAWYER_ID).text
    assert "Демо открыл новый посетитель" in note and "@ivan_lawyer" in note
    assert not (await ivan.send("/start")).to(LAWYER_ID), "о том же посетителе — один раз"


async def test_phone_not_asked(ivan):
    await ivan.send("/start")
    await ivan.press("consent")
    replies = await ivan.send("Иванов Иван Иванович")
    assert db.get_client(ivan.id)["phone"] == texts.DEMO_PHONE
    assert texts.DEMO_HINT in replies.text
    assert texts.BTN_PHONE not in replies.buttons


async def test_visitor_reviews_own_documents(ivan):
    await ivan.register()
    replies = await ivan.submit("passport")
    assert "Новые документы на проверку" in replies.to(ivan.id).text, "уведомление юриста — самому посетителю"
    assert not replies.to(LAWYER_ID), "владельцу чужие файлы не приходят"

    replies = await ivan.press(kb.ReviewCallback(action="accept", client_id=ivan.id, doc_id="passport"))
    assert db.get_status(ivan.id, "passport") == "accepted"


async def test_visitors_do_not_see_each_other(ivan, petr):
    await ivan.register()
    await ivan.submit("passport")
    await petr.register(full_name="Петров Пётр Петрович")

    replies = await petr.send("/clients")
    assert not any("Иванов" in button for button in replies.buttons)

    # Даже подделав данные кнопки, до чужого клиента не добраться
    for data in (kb.ClientCallback(action="open", client_id=ivan.id),
                 kb.ReviewCallback(action="files", client_id=ivan.id, doc_id="passport"),
                 kb.ReviewCallback(action="accept", client_id=ivan.id, doc_id="passport")):
        replies = await petr.press(data)
        assert texts.STALE_BUTTON in replies.alerts
        assert not replies.of("SendMessage") and not replies.of("SendMediaGroup")
    assert db.get_status(ivan.id, "passport") == "review"


async def test_fake_clients_shown_but_not_reminded(ivan):
    seed_demo.seed()
    await ivan.register()
    replies = await ivan.send("/clients")
    assert any("Смирнова Анна Викторовна" in button for button in replies.buttons)

    replies = await ivan.press(kb.ClientCallback(action="open", client_id=-2))
    assert "Кузнецов Дмитрий Олегович" in replies.text
    replies = await ivan.press(kb.ClientCallback(action="remind", client_id=-2))
    assert texts.DEMO_FAKE_REMIND in replies.alerts


async def test_export_contains_only_own_and_fake_clients(ivan, petr):
    seed_demo.seed()
    await ivan.register()
    await petr.register(full_name="Петров Пётр Петрович")
    replies = await ivan.send("/export")
    [document] = replies.of("SendDocument")
    from io import BytesIO
    from openpyxl import load_workbook
    names = [row[0] for row in load_workbook(BytesIO(document.document.data)).active.iter_rows(min_row=2, values_only=True)]
    assert "Иванов Иван Иванович" in names and "Смирнова Анна Викторовна" in names
    assert "Петров Пётр Петрович" not in names


async def test_reset_deletes_own_data(ivan):
    await ivan.register()
    await ivan.submit("passport")
    replies = await ivan.send("/reset")
    assert texts.DEMO_RESET_DONE in replies.text
    assert db.get_client(ivan.id) is None
    replies = await ivan.send("/start")
    assert "согласие" in replies.text, "можно пройти заново"


async def test_no_automatic_reminders(tg, ivan, daytime):
    await ivan.register()
    set_client(ivan.id, last_activity_at=LONG_AGO)
    assert not await tg.run(reminders.check_reminders(tg.bot))


async def test_manual_reminder_works_on_own_card(ivan):
    await ivan.register()
    replies = await ivan.press(kb.ClientCallback(action="remind", client_id=ivan.id))
    assert "Осталось" in replies.to(ivan.id).text


# ---------- вне демо ----------

async def test_reset_not_available_outside_demo(ivan, monkeypatch):
    monkeypatch.setattr(config, "DEMO_MODE", False)
    await ivan.register()
    replies = await ivan.send("/reset")
    assert texts.UNKNOWN in replies.text
    assert db.get_client(ivan.id) is not None


async def test_fake_clients_never_get_automatic_reminders(tg, monkeypatch, daytime):
    monkeypatch.setattr(config, "DEMO_MODE", False)
    seed_demo.seed()
    for number in range(1, len(seed_demo.CLIENTS) + 1):
        set_client(-number, last_activity_at=LONG_AGO)
    assert not await tg.run(reminders.check_reminders(tg.bot))


def test_seed_is_repeatable_and_clearable():
    assert seed_demo.seed() == len(seed_demo.CLIENTS)
    assert seed_demo.seed() == len(seed_demo.CLIENTS)  # повторный запуск не плодит дублей
    assert len([c for c in db.get_clients() if c["id"] < 0]) == len(seed_demo.CLIENTS)
    assert db.delete_demo_clients() == len(seed_demo.CLIENTS)
    assert db.get_clients() == []
