"""Картинки для README из настоящих ответов бота.

Сценарии прогоняются через ту же подставную обвязку Telegram, что и тесты (tests/conftest.py):
бот отвечает своими настоящими текстами и кнопками, а скрипт рисует переписку в стиле Telegram
и снимает её браузером Edge. Поменяли тексты в texts.py — перезапустите скрипт, картинки обновятся.
Это отрисовка переписки, а не снимок экрана приложения Telegram.

Запуск из PowerShell (нужны Microsoft Edge и pip install -r requirements-dev.txt):
    .venv\\Scripts\\python -m scripts.render_screenshots
Из Git Bash Edge не запускается — только из PowerShell или cmd.
"""
# tests.conftest — первым: он задаёт настройки и временную базу до импорта бота
from tests.conftest import LAWYER_ID, Chat  # isort: skip

import asyncio
import base64
import copy
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from html import escape
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from aiogram import Bot, Dispatcher
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.types import InlineKeyboardMarkup, InputMediaPhoto, ReplyKeyboardMarkup, ReplyKeyboardRemove
from openpyxl import load_workbook
from PIL import Image

from bot import config, db, texts
from bot.config import BASE_DIR
from bot.handlers import client, lawyer
from scripts import seed_demo
from tests.conftest import FakeTelegram

EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
OUT = BASE_DIR / "docs"
PLACEHOLDER = "⟦msg⟧"  # текст сообщения с кнопкой: бот дописывает к нему «Принято» — подставим настоящий
CHAT_W, CHAT_H = 400, 800
BOT_NAME = "Сбор документов. БФЛ"


# ---------- модель переписки ----------

@dataclass
class Item:
    side: str                                    # in — бот, out — пользователь
    html: str = ""
    buttons: list = field(default_factory=list)  # [[(текст, callback_data)]]
    media: list = field(default_factory=list)    # "photo" или ("pdf", имя)
    time: str = ""


class Studio:
    """Несколько чатов сразу: ответ бота раскладывается по чатам, куда он ушёл."""

    def __init__(self, tg):
        self.tg = tg
        self.screens: dict[int, "Screen"] = {}
        self.minute = 14

    def clock(self) -> str:
        self.minute += 1
        return f"10:{self.minute:02d}"

    def screen(self, chat: Chat, keyboard_visible: bool = True) -> "Screen":
        screen = Screen(self, chat, keyboard_visible)
        self.screens[chat.id] = screen
        return screen

    def absorb(self, replies, actor: "Screen") -> None:
        for request in replies:
            chat_id = getattr(request, "chat_id", None)
            screen = actor if chat_id is None else self.screens.get(chat_id)
            if screen:
                screen.apply(request)


class Screen:
    def __init__(self, studio: Studio, chat: Chat, keyboard_visible: bool):
        self.studio, self.chat = studio, chat
        self.items: list[Item] = []
        self.keyboard = None
        self.keyboard_visible = keyboard_visible
        self.edit_target: Item | None = None

    def _out(self, **fields) -> None:
        self.items.append(Item("out", time=self.studio.clock(), **fields))

    async def send(self, text: str) -> None:
        self._out(html=escape(text))
        self.studio.absorb(await self.chat.send(text), self)

    async def album(self, count: int) -> None:
        self._out(media=["photo"] * count)
        self.studio.absorb(await self.chat.send_album(count), self)

    async def pdf(self, name: str) -> None:
        self._out(media=[("pdf", name)])
        self.studio.absorb(await self.chat.send_pdf(name), self)

    async def contact(self, phone: str) -> None:
        self._out(html=f"📱 +{phone}")
        self.studio.absorb(await self.chat.send_contact(phone), self)

    async def press(self, button: str, in_message: str = "") -> None:
        """Нажать кнопку, чья подпись начинается с button, в последнем сообщении, где она есть
        (и где есть текст in_message)."""
        for item in reversed(self.items):
            if in_message not in item.html:
                continue
            for row in item.buttons:
                for text, data in row:
                    if text.startswith(button):
                        self.edit_target = item
                        self.studio.absorb(await self.chat.press(data, message_text=PLACEHOLDER), self)
                        return
        raise LookupError(f"нет кнопки «{button}»")

    def apply(self, request) -> None:
        name = type(request).__name__
        now = f"10:{self.studio.minute:02d}"
        if name == "SendMessage":
            item = Item("in", html=request.text, time=now)
            markup = request.reply_markup
            if isinstance(markup, InlineKeyboardMarkup):
                item.buttons = inline_rows(markup)
            elif isinstance(markup, ReplyKeyboardMarkup):
                self.keyboard = [[button.text for button in row] for row in markup.keyboard]
            elif isinstance(markup, ReplyKeyboardRemove):
                self.keyboard = None
            self.items.append(item)
        elif name == "SendMediaGroup":
            kinds = ["photo" if isinstance(m, InputMediaPhoto) else ("pdf", "скан.pdf") for m in request.media]
            self.items.append(Item("in", media=kinds, time=now))
        elif name == "SendPhoto":
            self.items.append(Item("in", media=["photo"], time=now))
        elif name == "SendDocument":
            filename = getattr(request.document, "filename", None) or "скан.pdf"
            self.items.append(Item("in", media=[("pdf", filename)], html=request.caption or "", time=now))
        elif name == "EditMessageText" and self.edit_target:
            self.edit_target.html = request.text.replace(PLACEHOLDER, self.edit_target.html)
            self.edit_target.buttons = inline_rows(request.reply_markup)
        elif name == "EditMessageReplyMarkup" and self.edit_target:
            self.edit_target.buttons = inline_rows(request.reply_markup)

    def mark(self) -> int:
        return len(self.items)

    def snapshot(self, start: int = 0) -> dict:
        return {"items": copy.deepcopy(self.items[start:]), "keyboard": copy.deepcopy(self.keyboard),
                "keyboard_visible": self.keyboard_visible}


def inline_rows(markup) -> list:
    if not isinstance(markup, InlineKeyboardMarkup):
        return []
    return [[(button.text, button.callback_data) for button in row] for row in markup.inline_keyboard]


# ---------- вёрстка в стиле Telegram ----------

CSS = """
* { box-sizing: border-box; margin: 0; padding: 0; }
body { width: %(w)dpx; height: %(h)dpx; overflow: hidden; display: flex; flex-direction: column;
       font: 15px/1.36 "Segoe UI", system-ui, sans-serif; color: #000; background: #fff; }
.header { height: 56px; display: flex; align-items: center; gap: 12px; padding: 0 14px;
          background: #fff; border-bottom: 1px solid #e3e6ea; flex-shrink: 0; }
.header .back { font-size: 22px; color: #8a96a3; }
.header img { width: 40px; height: 40px; border-radius: 50%%; }
.header b { display: block; font-size: 16px; }
.header span { font-size: 13px; color: #8a96a3; }
.chat { flex: 1; display: flex; flex-direction: column; justify-content: flex-end; overflow: hidden;
        gap: 6px; padding: 10px 10px 8px;
        background: linear-gradient(160deg, #c4d9b3 0%%, #a9cba6 45%%, #93bdb0 100%%); }
.msg { max-width: 84%%; }
.msg.in { align-self: flex-start; }
.msg.out { align-self: flex-end; }
.bubble { padding: 7px 10px 6px; border-radius: 14px; box-shadow: 0 1px 1px rgba(0,0,0,.12); word-wrap: break-word; }
.bubble::after { content: ""; display: block; clear: both; }  /* время не выпадает из пузыря */
.in .bubble { background: #fff; border-bottom-left-radius: 5px; }
.out .bubble { background: #effdde; border-bottom-right-radius: 5px; }
.time { float: right; margin: 4px 0 -2px 10px; font-size: 11.5px; color: #a0acb6; }
.out .time { color: #5dab4f; }
.buttons { margin-top: 3px; display: flex; flex-direction: column; gap: 3px; }
.row { display: flex; gap: 3px; }
.btn { flex: 1; text-align: center; padding: 8px 6px; border-radius: 9px; color: #fff; font-weight: 600;
       font-size: 14px; background: rgba(40, 70, 55, .34); }
.media { display: grid; gap: 2px; border-radius: 14px; overflow: hidden; box-shadow: 0 1px 1px rgba(0,0,0,.12); }
.media.n1 { grid-template-columns: 240px; }
.media.n2, .media.n4 { grid-template-columns: 124px 124px; }
.media.n3 { grid-template-columns: 124px 124px; }
.media.n3 svg:first-child { grid-column: span 2; width: 250px; height: 118px; }
.media svg { width: 124px; height: 92px; display: block; }
.media.n1 svg { width: 200px; height: 150px; }
.file { display: flex; gap: 10px; align-items: center; }
.file .icon { width: 44px; height: 44px; border-radius: 50%%; background: #4fae4e; color: #fff;
              display: flex; align-items: center; justify-content: center; font-size: 20px; }
.in .file .icon { background: #3a95d8; }
.file b { display: block; font-size: 14.5px; }
.file small { color: #7d8b97; font-size: 13px; }
.input { height: 50px; flex-shrink: 0; display: flex; align-items: center; gap: 14px; padding: 0 16px;
         background: #fff; border-top: 1px solid #e3e6ea; color: #9aa6b1; font-size: 16px; }
.input .grow { flex: 1; }
.keyboard { flex-shrink: 0; background: #f0f2f5; padding: 6px 6px 10px; display: flex; flex-direction: column; gap: 6px; }
.keyboard .row { gap: 6px; }
.key { flex: 1; text-align: center; background: #fff; border-radius: 9px; padding: 11px 4px;
       font-size: 14.5px; box-shadow: 0 1px 1px rgba(0,0,0,.15); }
"""


def doc_photo(seed: int) -> str:
    """Фото документа на столе — условная картинка вместо настоящего скана."""
    angle = (-3, 2, -1, 3)[seed % 4]
    widths = [(70, 50), (90, 40), (60, 70), (80, 55)][seed % 4]
    lines = "".join(
        f'<rect x="110" y="{40 + 14 * i}" width="{widths[i % 2]}" height="6" rx="3" fill="#b9c2cb"/>' for i in range(5)
    )
    return (
        '<svg viewBox="0 0 240 180" preserveAspectRatio="xMidYMid slice" xmlns="http://www.w3.org/2000/svg">'
        '<rect width="240" height="180" fill="#6d7a84"/>'
        f'<g transform="rotate({angle} 120 90)">'
        '<rect x="40" y="16" width="160" height="148" rx="5" fill="#f6f3ea"/>'
        '<rect x="54" y="34" width="44" height="56" rx="3" fill="#c8d4de"/>'
        f'{lines}'
        '<rect x="54" y="110" width="130" height="6" rx="3" fill="#d3d9df"/>'
        '<rect x="54" y="124" width="100" height="6" rx="3" fill="#d3d9df"/>'
        '<rect x="54" y="138" width="116" height="6" rx="3" fill="#d3d9df"/>'
        '</g></svg>'
    )


def render_item(item: Item, seed: int) -> str:
    time = f'<span class="time">{item.time}{" ✓✓" if item.side == "out" else ""}</span>'
    parts = []
    photos = [m for m in item.media if m == "photo"]
    files = [m for m in item.media if m != "photo"]
    if photos:
        parts.append(f'<div class="media n{min(len(photos), 4)}">'
                     + "".join(doc_photo(seed + i) for i in range(len(photos))) + "</div>")
    for _, name in files:
        size = "84 КБ · XLSX" if name.endswith(".xlsx") else "1,2 МБ · PDF"
        caption = f"<br>{item.html.replace(chr(10), '<br>')}" if item.html else ""
        parts.append(f'<div class="bubble"><div class="file"><div class="icon">📄</div>'
                     f'<div><b>{escape(name)}</b><small>{size}</small></div></div>{caption}{time}</div>')
    if item.html and not files:
        parts.append(f'<div class="bubble">{item.html.replace(chr(10), "<br>")}{time}</div>')
    if item.buttons:
        rows = "".join('<div class="row">' + "".join(f'<div class="btn">{escape(text)}</div>' for text, _ in row)
                       + "</div>" for row in item.buttons)
        parts.append(f'<div class="buttons">{rows}</div>')
    return f'<div class="msg {item.side}">{"".join(parts)}</div>'


def chat_page(shot: dict) -> str:
    avatar = base64.b64encode((OUT / "avatar.png").read_bytes()).decode()
    items = "".join(render_item(item, seed) for seed, item in enumerate(shot["items"]))
    keyboard = ""
    if shot["keyboard"] and shot["keyboard_visible"]:
        keyboard = '<div class="keyboard">' + "".join(
            '<div class="row">' + "".join(f'<div class="key">{escape(key)}</div>' for key in row) + "</div>"
            for row in shot["keyboard"]) + "</div>"
    return f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8">
<style>{CSS % {"w": CHAT_W, "h": CHAT_H}}</style></head><body>
<div class="header"><span class="back">←</span><img src="data:image/png;base64,{avatar}">
<div><b>{BOT_NAME}</b><span>бот</span></div></div>
<div class="chat">{items}</div>
<div class="input"><span>📎</span><span class="grow">Сообщение</span><span>⌨</span></div>
{keyboard}
</body></html>"""


def excel_page(xlsx: bytes, filename: str) -> tuple[str, int, int]:
    """Выгрузка в Excel — таблицей в стиле Excel из настоящего файла, который прислал бот."""
    sheet = load_workbook(BytesIO(xlsx)).active
    columns = 11  # A–K: остальное «за краем окна», как при прокрутке
    widths = [int((sheet.column_dimensions[chr(65 + i)].width or 10) * 7.2 + 8) for i in range(columns)]
    head = '<tr><th class="corner"></th>' + "".join(
        f'<th style="width:{w}px">{chr(65 + i)}</th>' for i, w in enumerate(widths)) + "</tr>"
    body = ""
    note_shown = False
    for r, row in enumerate(sheet.iter_rows(min_row=1, max_col=columns), start=1):
        cells = ""
        for cell in row:
            style = ""
            rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            if isinstance(rgb, str):
                style = f"background:#{rgb[-6:]};"
            classes = ["hdr"] if r == 1 else []
            if cell.column == 1:
                classes.append("frozen")
            extra = ""
            if cell.comment:
                classes.append("commented")
                if not note_shown:
                    extra = f'<div class="note"><b>{escape(cell.comment.author)}:</b><br>{escape(cell.comment.text)}</div>'
                    note_shown = True
            cells += f'<td class="{" ".join(classes)}" style="{style}">{escape(str(cell.value or ""))}{extra}</td>'
        body += f'<tr class="{"top" if r == 1 else ""}"><th>{r}</th>{cells}</tr>'
    width, height = 1400, 128 + 48 + 25 * (sheet.max_row - 1) + 60
    page = f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"><style>
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ width: {width}px; height: {height}px; overflow: hidden; font: 14px Calibri, "Segoe UI", sans-serif; background: #fff; }}
.title {{ height: 34px; background: #217346; color: #fff; display: flex; align-items: center; justify-content: center; font-size: 13px; }}
.ribbon {{ height: 58px; background: #f3f3f3; border-bottom: 1px solid #d5d5d5; display: flex; align-items: center; gap: 26px;
           padding: 0 18px; color: #444; font-size: 13px; }}
.ribbon b {{ color: #217346; border-bottom: 3px solid #217346; padding-bottom: 4px; }}
.fx {{ height: 32px; border-bottom: 1px solid #d5d5d5; display: flex; align-items: center; font-size: 13px; color: #333; }}
.fx .ref {{ width: 90px; padding-left: 10px; border-right: 1px solid #d5d5d5; }}
.fx .f {{ padding: 0 14px; color: #888; font-style: italic; }}
table {{ border-collapse: collapse; table-layout: fixed; }}
th {{ background: #f3f3f3; color: #555; font-weight: 400; font-size: 12px; border: 1px solid #d5d5d5; height: 22px; }}
th.corner, tr th:first-child {{ width: 40px; }}
td {{ border: 1px solid #dadada; height: 25px; padding: 0 6px; white-space: nowrap; overflow: visible; position: relative; }}
td.hdr {{ font-weight: 700; white-space: normal; vertical-align: top; height: 48px; padding-top: 3px; line-height: 1.2; overflow: hidden; }}
tr.top td {{ border-bottom: 2px solid #9ab09f; }}
td.frozen {{ border-right: 2px solid #9ab09f; }}
td.commented::after {{ content: ""; position: absolute; top: 0; right: 0; border-left: 7px solid transparent; border-top: 7px solid #d32f2f; }}
.note {{ position: absolute; left: calc(100% + 14px); top: -4px; z-index: 5; width: 230px; white-space: normal;
         background: #fffbe0; border: 1px solid #b9ad62; box-shadow: 2px 2px 4px rgba(0,0,0,.2); padding: 6px 8px;
         font-size: 13px; line-height: 1.3; }}
.tabs {{ position: absolute; bottom: 0; left: 0; right: 0; height: 30px; background: #f3f3f3; border-top: 1px solid #d5d5d5;
         display: flex; align-items: stretch; padding-left: 60px; font-size: 13px; }}
.tabs span {{ background: #fff; padding: 6px 18px; color: #217346; font-weight: 600; border-bottom: 3px solid #217346; }}
</style></head><body>
<div class="title">{escape(filename)} — Excel</div>
<div class="ribbon"><span>Файл</span><b>Главная</b><span>Вставка</span><span>Разметка страницы</span><span>Формулы</span><span>Данные</span><span>Рецензирование</span><span>Вид</span></div>
<div class="fx"><span class="ref">A1</span><span class="f">fx</span><span>Клиент</span></div>
<table>{head}{body}</table>
<div class="tabs"><span>{escape(texts.EXPORT_SHEET)}</span></div>
</body></html>"""
    return page, width, height


# ---------- съёмка ----------

def shoot(html: str, png: Path, width: int, height: int, scale: float = 2) -> None:
    tmp = Path(tempfile.mkdtemp())
    page = tmp / "page.html"
    page.write_text(html, encoding="utf-8")
    shot = tmp / "shot.png"
    subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                    f"--force-device-scale-factor={scale}", f"--user-data-dir={tmp / 'profile'}",
                    f"--window-size={width},{height}", f"--screenshot={shot}", page.as_uri()],
                   capture_output=True, timeout=90)
    if not shot.exists():
        raise SystemExit("Edge не сделал снимок. Запускайте скрипт из PowerShell или cmd, не из Git Bash.")
    png.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(shot, png)
    shutil.rmtree(tmp, ignore_errors=True)
    print("  ✓", png.name)


def make_gif(frames: list[tuple[dict, int]], path: Path) -> None:
    tmp = Path(tempfile.mkdtemp())
    images = []
    for number, (shot, _) in enumerate(frames):
        png = tmp / f"{number}.png"
        shoot(chat_page(shot), png, CHAT_W, CHAT_H, scale=1.5)
        images.append(Image.open(png).convert("RGB").quantize(colors=160, dither=Image.Dither.NONE))
    images[0].save(path, save_all=True, append_images=images[1:], duration=[d for _, d in frames], loop=0, optimize=True)
    print(f"  ✓ {path.relative_to(BASE_DIR)} — {path.stat().st_size // 1024} КБ, кадров: {len(images)}")


# ---------- сценарии ----------

def make_tg() -> SimpleNamespace:
    api = FakeTelegram()
    bot = Bot("123456:TEST", session=api, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp = Dispatcher()
    dp.include_routers(lawyer.router, client.router)
    return SimpleNamespace(api=api, bot=bot, dp=dp)


async def register(screen: Screen) -> None:
    await screen.send("/start")
    await screen.press(texts.BTN_CONSENT)
    await screen.send("Иванов Иван Иванович")
    if not config.DEMO_MODE:
        await screen.contact("79991234567")


async def product_screens(tg) -> dict:
    """Настоящий режим: клиент и юрист — разные чаты."""
    studio = Studio(tg)
    ivan = studio.screen(Chat(tg, 111, "Иван"))
    lawyer_screen = studio.screen(Chat(tg, LAWYER_ID, "Юрист"), keyboard_visible=False)
    shots = {}

    await register(ivan)

    start = ivan.mark()
    await ivan.send(texts.BTN_SEND)
    await ivan.press("⬜ Паспорт")
    await ivan.album(3)
    await ivan.press(texts.BTN_DONE, in_message="Получено")
    shots["02-client-upload"] = ivan.snapshot(start)

    await lawyer_screen.press(texts.BTN_FILES, in_message="Паспорт")
    shots["04-lawyer-notification"] = lawyer_screen.snapshot()

    await ivan.send(texts.BTN_SEND)
    await ivan.press("⬜ СНИЛС")
    await ivan.album(1)
    await ivan.press(texts.BTN_DONE, in_message="Получено")
    await ivan.send(texts.BTN_SEND)
    await ivan.press("⬜ Свидетельство о браке")
    await ivan.press(texts.BTN_NOT_APPLICABLE)
    await ivan.send(texts.BTN_SEND)
    await ivan.press("⬜ ИНН")
    await ivan.pdf("ИНН скан.pdf")
    await ivan.press(texts.BTN_DONE, in_message="Получено")

    await lawyer_screen.press(texts.BTN_ACCEPT, in_message="Паспорт")
    await lawyer_screen.press(texts.BTN_REJECT, in_message="Документ: <b>ИНН")
    await lawyer_screen.send("Скан обрезан снизу — не видно номер. Пришлите документ целиком")
    shots["03-client-returned"] = ivan.snapshot()

    await ivan.send(texts.BTN_STATUS)
    shots["01-client-checklist"] = ivan.snapshot()

    seed_demo.seed()
    await lawyer_screen.send("/clients")
    shots["05-lawyer-clients"] = lawyer_screen.snapshot()

    await lawyer_screen.press("Иванов Иван Иванович")
    shots["06-lawyer-card"] = lawyer_screen.snapshot()

    replies = await lawyer_screen.chat.send("/export")
    [document] = replies.of("SendDocument")
    shots["07-excel"] = (document.document.data, document.document.filename)
    return shots


async def demo_gif_frames(tg) -> list[tuple[dict, int]]:
    """Демо-режим: посетитель — и клиент, и юрист в одном чате. Кадры для GIF в шапке README."""
    config.DEMO_MODE = True
    studio = Studio(tg)
    visitor = studio.screen(Chat(tg, 333, "Посетитель"))
    await register(visitor)
    frames = []

    def frame(duration: int) -> None:
        frames.append((visitor.snapshot(), duration))

    frame(1600)
    await visitor.send(texts.BTN_SEND)
    frame(1400)
    await visitor.press("⬜ Паспорт")
    frame(1400)
    await visitor.album(3)
    frame(1400)
    await visitor.press(texts.BTN_DONE, in_message="Получено")
    frame(2200)
    await visitor.press(texts.BTN_ACCEPT, in_message="Новые документы")
    frame(2400)
    await visitor.send(texts.BTN_STATUS)
    frame(3600)
    config.DEMO_MODE = False
    return frames


async def main() -> None:
    db.init_db()
    tg = make_tg()
    print("Сценарии…")
    shots = await product_screens(tg)
    frames = await demo_gif_frames(tg)

    print("Снимки:")
    for name, shot in shots.items():
        if name == "07-excel":
            page, width, height = excel_page(*shot)
            shoot(page, OUT / "screenshots" / f"{name}.png", width, height, scale=1.5)
        else:
            shoot(chat_page(shot), OUT / "screenshots" / f"{name}.png", CHAT_W, CHAT_H)
    make_gif(frames, OUT / "demo.gif")


if __name__ == "__main__":
    asyncio.run(main())
